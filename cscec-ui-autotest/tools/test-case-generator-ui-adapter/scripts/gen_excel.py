#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
通用Excel测试用例生成器 - 从Markdown表格生成Excel文件

用法（Windows）:
    py -3 gen_excel.py --input "用例.md" --output "输出.xlsx"

用法（已配置 python 命令时）:
    python gen_excel.py --input "用例.md" --output "输出.xlsx"

Markdown格式要求:
    表头固定为: | 所属模块 | 标题 | 类型 | 优先级 | 前置条件 | 负责人 | 步骤 | 预期结果 | 所涉及到的数据 |
    支持 GFM 表格格式

也支持直接传入JSON数据文件:
    py -3 gen_excel.py --json cases.json --output "输出.xlsx"
"""

import argparse
import json
import os
import re
import sys
from html import unescape

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误：缺少 openpyxl 库，请先安装: pip install openpyxl")
    sys.exit(1)


# 表头定义
HEADERS = ["所属模块", "标题", "类型", "优先级", "前置条件", "负责人", "步骤", "预期结果", "所涉及到的数据"]

HEADER_ALIASES = [
    {"所属模块", "模块"},
    {"标题"},
    {"类型"},
    {"优先级"},
    {"前置条件"},
    {"负责人"},
    {"步骤"},
    {"预期结果"},
    {"所涉及到的数据", "涉及到的数据", "涉及的数据"},
]

TYPE_SPLIT_RE = re.compile(r"[/,，、]")

# 列宽配置
COL_WIDTHS = [16, 42, 12, 8, 40, 10, 45, 40, 35]

# 样式定义
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_FONT = Font(size=10)
DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 优先级样式
PRIORITY_STYLES = {
    "P0": {
        "fill": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "font": Font(size=10, bold=True, color="9C0006")
    },
    "P1": {
        "fill": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "font": Font(size=10, bold=True, color="9C6500")
    },
    "P2": {
        "fill": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "font": Font(size=10, bold=True, color="006100")
    }
}


def parse_markdown_table(md_content):
    """
    解析Markdown表格为二维数组
    
    支持多表格拼接：Markdown文件中可能有多个独立表格（如M1~M9各一个），
    每个表格都有自己的表头行和分隔行，需要全部跳过。
    
    返回: [[col1, col2, ...], ...]
    """
    lines = md_content.strip().split('\n')
    rows = []
    
    for line in lines:
        stripped = line.strip()
        
        # 检测表格行，兼容全角分隔符
        if not stripped.startswith('|'):
            if not stripped.startswith('｜'):
                continue
        stripped = stripped.replace('｜', '|')
        
        # 分割单元格
        cells = [c.strip() for c in stripped.split('|')]
        # 过滤空字符串
        cells = [c for c in cells if c != '']
        
        # 跳过分隔行 (如 |---|---|)
        if all(re.match(r'^[-:]+$', c) for c in cells):
            continue
        
        # 跳过表头行，兼容历史字段名差异
        if len(cells) == len(HEADER_ALIASES):
            normalized_cells = [c.replace(' ', '') for c in cells]
            if all(normalized_cells[idx] in HEADER_ALIASES[idx] for idx in range(len(normalized_cells))):
                continue
        
        # 如果单元格数量与表头一致，添加为数据行
        if len(cells) == len(HEADERS):
            rows.append(cells)
        elif len(cells) > 0:
            # 尝试补齐或截断
            while len(cells) < len(HEADERS):
                cells.append("-")
            rows.append(cells[:len(HEADERS)])
    
    return rows


def _strip_html_tags(text):
    text = re.sub(r'(?i)<br\s*/?>', '\n', text)
    text = re.sub(r'<[^>]+>', '', text)
    return unescape(text).strip()


def parse_html_table(content):
    rows = []
    table_match = re.search(r'(?is)<table.*?</table>', content)
    if not table_match:
        return rows

    table_html = table_match.group(0)
    for tr in re.findall(r'(?is)<tr[^>]*>(.*?)</tr>', table_html):
        cells = re.findall(r'(?is)<t[hd][^>]*>(.*?)</t[hd]>', tr)
        if not cells:
            continue
        normalized = [_strip_html_tags(c) for c in cells]

        if len(normalized) == len(HEADER_ALIASES):
            compact = [c.replace(' ', '') for c in normalized]
            if all(compact[idx] in HEADER_ALIASES[idx] for idx in range(len(compact))):
                continue

        if len(normalized) == len(HEADERS):
            rows.append(normalized)
        elif len(normalized) > 0:
            while len(normalized) < len(HEADERS):
                normalized.append("-")
            rows.append(normalized[:len(HEADERS)])

    return rows


def normalize_multiline_text(text):
    text = str(text).replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'(?i)<br\s*/?>', '\n', text)
    text = re.sub(r'(?m)(^|\n)\s*(\d+)\.\s*', lambda m: f"{m.group(1)}【{m.group(2)}】", text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def parse_json_cases(json_path):
    """从JSON文件解析用例数据"""
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    if isinstance(data, list):
        return [item if isinstance(item, list) else list(item.values()) for item in data]
    elif isinstance(data, dict):
        cases = data.get('test_cases', data.get('cases', []))
        return [item if isinstance(item, list) else list(item.values()) for item in cases]
    else:
        raise ValueError("无法解析JSON数据格式")


def create_excel(cases_data, output_path, project_name="测试项目"):
    """
    从用例数据创建Excel文件
    
    cases_data: 二维数组，每行一条用例，列与HEADERS对应
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例"
    
    # 写入表头
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    
    # 写入数据
    for row_idx, case in enumerate(cases_data, 2):
        for col_idx, value in enumerate(case, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else "-")
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGNMENT
            cell.border = THIN_BORDER
            
            # 优先级颜色标记
            if col_idx == 4:  # 优先级列
                priority = str(value).strip().upper()
                if priority in PRIORITY_STYLES:
                    style = PRIORITY_STYLES[priority]
                    cell.fill = style["fill"]
                    cell.font = style["font"]
    
    # 设置列宽
    for i, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # 冻结首行
    ws.freeze_panes = "A2"
    
    # 设置自动筛选
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(cases_data)+1}"
    
    # 添加统计Sheet
    _add_stats_sheet(wb, cases_data)
    
    # 确保输出目录存在
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    wb.save(output_path)
    print(f"Excel文件已生成: {output_path}")
    print(f"总用例数: {len(cases_data)}")
    return output_path


def _add_stats_sheet(wb, cases_data):
    """添加用例统计Sheet"""
    ws2 = wb.create_sheet("用例统计")
    
    # 统计数据
    module_count = {}
    priority_count = {}
    type_count = {}
    
    for case in cases_data:
        module = str(case[0]).split("-")[0] if len(case) > 0 else "未知"
        module_count[module] = module_count.get(module, 0) + 1
        
        priority = str(case[3]) if len(case) > 3 else "未知"
        priority_count[priority] = priority_count.get(priority, 0) + 1
        
        types_str = str(case[2]) if len(case) > 2 else ""
        for t in TYPE_SPLIT_RE.split(types_str):
            t = t.strip()
            if t:
                type_count[t] = type_count.get(t, 0) + 1
    
    # 写入统计表头
    stats_headers = ["统计维度", "分类", "数量"]
    for col_idx, header in enumerate(stats_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    
    row = 2
    
    for module, count in sorted(module_count.items()):
        for col in [1, 2, 3]:
            ws2.cell(row=row, column=col).border = THIN_BORDER
        ws2.cell(row=row, column=1, value="模块")
        ws2.cell(row=row, column=2, value=module)
        ws2.cell(row=row, column=3, value=count)
        row += 1
    
    for priority, count in sorted(priority_count.items()):
        for col in [1, 2, 3]:
            ws2.cell(row=row, column=col).border = THIN_BORDER
        ws2.cell(row=row, column=1, value="优先级")
        ws2.cell(row=row, column=2, value=priority)
        ws2.cell(row=row, column=3, value=count)
        row += 1
    
    for t, count in sorted(type_count.items()):
        for col in [1, 2, 3]:
            ws2.cell(row=row, column=col).border = THIN_BORDER
        ws2.cell(row=row, column=1, value="类型")
        ws2.cell(row=row, column=2, value=t)
        ws2.cell(row=row, column=3, value=count)
        row += 1
    
    for col in [1, 2, 3]:
        ws2.cell(row=row, column=col).border = THIN_BORDER
    ws2.cell(row=row, column=1, value="总计")
    ws2.cell(row=row, column=1).font = Font(bold=True, size=11)
    ws2.cell(row=row, column=2, value="全部用例")
    ws2.cell(row=row, column=3, value=len(cases_data))
    ws2.cell(row=row, column=3).font = Font(bold=True, size=11, color="FF0000")
    
    ws2.column_dimensions['A'].width = 15
    ws2.column_dimensions['B'].width = 25
    ws2.column_dimensions['C'].width = 10


def main():
    parser = argparse.ArgumentParser(description='从Markdown/JSON生成Excel测试用例')
    parser.add_argument('--input', '-i', help='输入的Markdown文件路径')
    parser.add_argument('--json', '-j', help='输入的JSON数据文件路径')
    parser.add_argument('--output', '-o', required=True, help='输出的Excel文件路径')
    parser.add_argument('--project', '-p', default='测试项目', help='项目名称（默认：测试项目）')
    
    args = parser.parse_args()
    
    if not args.input and not args.json:
        print("错误：必须指定 --input 或 --json 参数")
        sys.exit(1)
    
    if args.json:
        cases_data = parse_json_cases(args.json)
    else:
        if not os.path.exists(args.input):
            print(f"错误：输入文件不存在: {args.input}")
            sys.exit(1)
        
        with open(args.input, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        if '<table' in md_content.lower():
            cases_data = parse_html_table(md_content)
        else:
            cases_data = parse_markdown_table(md_content)
    
    if not cases_data:
        print("警告：未解析到有效的用例数据")
        sys.exit(1)

    cases_data = [
        [normalize_multiline_text(col) for col in row]
        for row in cases_data
    ]
    
    create_excel(cases_data, args.output, args.project)


if __name__ == '__main__':
    main()
